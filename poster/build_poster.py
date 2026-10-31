"""Build the two-slide childhood poster. Text is read verbatim from the user's .docx."""
import sys, copy
import docx
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
DOC = docx.Document(SRC)
P = [p.text for p in DOC.paragraphs]

# ---- the only text edits (all reported to the user) ----
FIXES = {
    39: [("ecosystem (work of parents", "exosystem (work of parents")],
    20: [("development: Erikson.", "development: Erikson")],
    84: [("2024 to 2025, england", "2024 to 2025, England")],
}
def T(i):
    t = P[i].strip()
    for a, b in FIXES.get(i, []):
        assert a in t, (i, a)
        t = t.replace(a, b)
    return t

FONT = "Arial"
BODY = 9.0
SW, SH = 18.0, 10.125
LH = 1.2  # line height factor used for measurement
_fonts = {}
def font(bold, italic, size):
    k = (bold, italic, size)
    if k not in _fonts:
        name = "LiberationSans-" + ({(0, 0): "Regular", (1, 0): "Bold", (0, 1): "Italic", (1, 1): "BoldItalic"}[(int(bold), int(italic))])
        _fonts[k] = ImageFont.truetype(f"/usr/share/fonts/truetype/liberation/{name}.ttf", int(size * 20))
    return _fonts[k]

def text_w(s, bold=False, italic=False, size=BODY):
    return font(bold, italic, size).getlength(s) / 20.0  # points

def count_lines(runs, width_pt, size):
    """runs: list of (text, bold, italic). Greedy word wrap."""
    words = []
    for t, b, it in runs:
        for k, w in enumerate(t.split(" ")):
            words.append((w, b, it, k > 0))
    lines, cur = 1, 0.0
    for w, b, it, _ in words:
        ww = text_w(w, b, it, size)
        sp = text_w(" ", b, it, size)
        if cur == 0:
            cur = ww
        elif cur + sp + ww <= width_pt:
            cur += sp + ww
        else:
            lines += 1
            cur = ww
    return lines

# ---------------- colours ----------------
def rgb(h): return RGBColor.from_string(h)
INK = "1F2937"
COL = {
    "intro": "3D3B8E", "bio": "16796F", "psych": "6A3F9A",
    "socio": "B4462B", "key": "245C8F", "concl": "3D3B8E",
}
TINT = {
    "intro": "EEEDFA", "bio": "E4F4F1", "psych": "F1EAF8",
    "socio": "FBEBE6", "key": "E5EEF7", "concl": "EEEDFA",
}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK = prs.slide_layouts[6]

def set_bg_gradient(slide, c1, c2):
    fill = slide.background.fill
    fill.gradient()
    fill.gradient_angle = 90
    st = fill.gradient_stops
    st[0].color.rgb = rgb(c1); st[0].position = 0
    st[1].color.rgb = rgb(c2); st[1].position = 1.0

def box(slide, x, y, w, h, fillc=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fillc: s.fill.solid(); s.fill.fore_color.rgb = rgb(fillc)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if name: s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.08, 0.06 / max(min(w, h), 0.01) * 1.5)
    tf = s.text_frame
    return s

def set_bullet(p, indent_in=0.13, char="•", color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    if color:
        bc = etree.SubElement(pPr, qn("a:buClr")); sc = etree.SubElement(bc, qn("a:srgbClr")); sc.set("val", color)
    bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", char)

def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, qn("a:buNone"))

def add_runs(p, runs, size, color=INK):
    for t, b, it, *rest in runs:
        r = p.add_run(); r.text = t
        f = r.font; f.name = FONT; f.size = Pt(size); f.bold = b; f.italic = it
        f.color.rgb = rgb(rest[0] if rest and rest[0] else color)
        if len(rest) > 1 and rest[1]: f.underline = True

def textbox(slide, x, y, w, h, margin=(0.06, 0.03), anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin[0])
    tf.margin_top = tf.margin_bottom = Inches(margin[1])
    tf.vertical_anchor = anchor
    tf.auto_size = None
    if name: tb.name = name
    return tb, tf

def highlight(text, terms, color):
    """Split text into runs; terms are bolded in accent colour. Text unchanged."""
    runs = [(text, False, False, None)]
    for term in terms:
        out = []
        for t, b, it, c in runs:
            if b or term not in t:
                out.append((t, b, it, c)); continue
            parts = t.split(term)
            for k, part in enumerate(parts):
                if part: out.append((part, False, False, None))
                if k < len(parts) - 1: out.append((term, True, False, color))
        runs = out
    return runs

# ---------------- content blocks ----------------
# block = list of (kind, index, terms)   kind: sub | bullet | para
def block_items(spec, color):
    items = []
    for kind, i, terms in spec:
        t = T(i)
        if kind == "sub":
            items.append(("sub", [(t, True, False, color)]))
        else:
            items.append((kind, highlight(t, terms, color)))
    return items

SPACE_AFTER = 2.0
SUB_BEFORE = 3.0
BUL_IND = 0.13
def measure_items(items, w_in, size=BODY):
    inner = (w_in - 0.12) * 72
    h = 0.06 * 72
    for k, (kind, runs) in enumerate(items):
        width = inner - (BUL_IND * 72 if kind == "bullet" else 0)
        rr = [(t, b, it) for t, b, it, *_ in runs]
        sz = size + 0.5 if kind == "sub" else size
        h += count_lines(rr, width, sz) * sz * LH + SPACE_AFTER
        if kind == "sub" and k > 0: h += SUB_BEFORE
    return h / 72.0

def render_items(slide, items, x, y, w, h, color, size=BODY, name=None):
    tb, tf = textbox(slide, x, y, w, h, name=name)
    first = True
    for k, (kind, runs) in enumerate(items):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(SPACE_AFTER)
        p.line_spacing = 1.0
        if kind == "sub" and k > 0: p.space_before = Pt(SUB_BEFORE)
        if kind == "bullet": set_bullet(p, BUL_IND, color=color)
        add_runs(p, runs, size + 0.5 if kind == "sub" else size)
    return tb

HDR_H = 0.32
def panel(slide, x, y, w, h, key, title):
    box(slide, x, y, w, h, fillc="FFFFFF", line=COL[key], lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Panel {title}")
    hb = box(slide, x, y, w, HDR_H, fillc=COL[key], shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Header {title}")
    tf = hb.text_frame; tf.margin_left = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    add_runs(p, [(title, True, False, "FFFFFF")], 12.5)

def caption(slide, x, y, w, fig, cite, color):
    tb, tf = textbox(slide, x, y, w, 0.2, margin=(0.04, 0.0))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
    add_runs(p, [(f"Figure {fig}. ", True, False, color), ("[MISSING: caption for Figure %d] " % fig, False, True, "B42318"), (cite, False, False, "4B5563")], 8.5)
    return 0.2

def arrow(slide, x, y, w, h, direction, color):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW}[direction]
    a = box(slide, x, y, w, h, fillc=color, shape=shp, name=f"Flow arrow {direction}")
    return a

def label(slide, x, y, w, h, text, size=8.5, bold=False, color=INK, fillc=None, line=None, align=PP_ALIGN.CENTER,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, italic=False, margin=0.03):
    s = box(slide, x, y, w, h, fillc=fillc, line=line, lw=0.75, shape=shape)
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for k, ln in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        add_runs(p, [(ln, bold if k == 0 else False, italic, color)], size)
    return s

# ======================= SLIDE 1 =======================
s1 = prs.slides.add_slide(BLANK)
set_bg_gradient(s1, "F3EEFB", "E3F4F0")

M = 0.35; G = 0.32; TOP = 1.25; BOT = SH - 0.3
NC = 5
CW = (SW - 2 * M - (NC - 1) * G) / NC
cx = [M + k * (CW + G) for k in range(NC)]

# Title banner
ban = box(s1, M, 0.25, SW - 2 * M, 0.85, fillc="2E2A6B", shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Title banner")
tf = ban.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
add_runs(p, [("Human Development in Childhood", True, False, "FFFFFF"), ("  –  [Student ID]", True, False, "FFD166")], 30)

# ---------- section specs (paragraph indices from the .docx) ----------
intro = [("para", 1, ["early childhood (2 to 6 years)", "middle childhood (6 to 12 years)", "biopsychosocial perspective"])]
bio = [("sub", 3, []), ("bullet", 4, []), ("bullet", 5, []),
       ("sub", 6, []), ("bullet", 7, ["Gross motor skills:"]), ("bullet", 8, ["Fine motor skills:"]),
       ("sub", 9, []), ("bullet", 10, ["memory, attention and self-control"]),
       ("sub", 11, []), ("bullet", 12, []), ("bullet", 13, ["81%"]), ("bullet", 14, [])]
psych_a = [("sub", 16, []), ("bullet", 17, ["Preoperational stage (2-7 years):"]),
           ("bullet", 18, ["Concrete operational stage (7-11 years):"]), ("bullet", 19, ["Critique:"]),
           ("sub", 20, []), ("bullet", 21, ["Initiative vs guilt (3 to 6 years):"]),
           ("bullet", 22, ["Industry vs. inferiority (6 to 12 years):"]), ("bullet", 23, ["theory of mind"])]
psych_b = [("sub", 24, []), ("bullet", 25, ["thinking", "emotions and social life"]), ("bullet", 26, ["Mutual principles:"])]
socio = [("sub", 28, []), ("bullet", 29, []), ("bullet", 30, ["secure attachment"]),
         ("sub", 31, []), ("bullet", 32, []), ("bullet", 33, []),
         ("sub", 34, []), ("bullet", 35, []), ("bullet", 36, ["trust and loyalty"])]
key_a = [("sub", 38, []), ("bullet", 39, ["microsystem", "mesosystem", "exosystem", "macrosystem", "chronosystem"]),
         ("bullet", 40, ["process, person, context and time"])]
key_b = [("sub", 41, []), ("bullet", 42, ["27%", "1.9 million children (13%)"]), ("bullet", 43, ["family income"])]
key_c = [("sub", 44, []), ("bullet", 45, ["22.2%"]), ("bullet", 46, ["29.3%", "13.5%"])]
key_d = [("sub", 47, []), ("bullet", 48, []), ("bullet", 49, ["four or more ACEs"]),
         ("sub", 50, []), ("para", 51, ["interdependence"])]
concl = [("bullet", 53, []), ("bullet", 54, ["interrelated"]), ("bullet", 55, ["active", "environment does matter"]),
         ("bullet", 56, ["caring relationships"]), ("bullet", 57, ["Key message:"])]

PAD = 0.07  # inner padding of panels
def IW(ncols=1): return CW * ncols + G * (ncols - 1) - 2 * PAD

report = []
def place_text(spec, key, x, y, w, name):
    items = block_items(spec, COL[key])
    h = measure_items(items, w)
    render_items(s1, items, x, y, w, h, COL[key], name=name)
    return h

# ---------------- visual builders ----------------
def v_timeline(x, y, w):
    """Figure 1: childhood stages and motor / brain progressions (Boyd and Bee, 2019)."""
    c = COL["bio"]; h0 = y
    lab = [("Infancy", 0.2, "D1D5DB", INK), ("Early childhood\n2 to 6 years", 0.3, "8FD3C7", INK),
           ("Middle childhood\n6 to 12 years", 0.3, c, "FFFFFF"), ("Adolescence", 0.2, "D1D5DB", INK)]
    xx = x
    for t, frac, fc, tc in lab:
        ww = w * frac
        label(s1, xx, y, ww - 0.02, 0.38, t, size=7, bold=True, color=tc, fillc=fc, shape=MSO_SHAPE.RECTANGLE)
        xx += ww
    y += 0.46
    rows = [("Gross motor", ["running", "hopping", "climbing", "riding a bike"]),
            ("Fine motor", ["holding a pencil", "using scissors", "joined-up writing"]),
            ("Brain", ["connections among nerve cells", "memory, attention and self-control"])]
    lw = 0.62
    for name, steps in rows:
        label(s1, x, y, lw, 0.38, name, size=7.5, bold=True, color="FFFFFF", fillc=c)
        n = len(steps); aw = 0.1
        sw = (w - lw - 0.06 - (n - 1) * (aw + 0.04)) / n
        xx = x + lw + 0.06
        for k, st in enumerate(steps):
            label(s1, xx, y, sw, 0.38, st, size=7, color=INK, fillc=TINT["bio"], line=c)
            xx += sw
            if k < n - 1:
                box(s1, xx + 0.02, y + 0.13, aw, 0.12, fillc=c, shape=MSO_SHAPE.RIGHT_ARROW)
                xx += aw + 0.04
        y += 0.44
    return y - h0

def v_stages(x, y, w):
    """Figure 2: Piaget vs Erikson stage comparison on an age axis 2-12."""
    c = COL["psych"]; h0 = y
    lw = 0.68; ax0 = x + lw + 0.05; ax1 = x + w - 0.05
    def ax(a): return ax0 + (a - 2) / 10.0 * (ax1 - ax0)
    rows = [("Piaget\nthinking", [("Preoperational\n(2-7 years)", 2, 7, "C9B6E4", INK), ("Concrete operational\n(7-11 years)", 7, 11, c, "FFFFFF")]),
            ("Erikson\nemotions and social life", [("Initiative vs guilt\n(3 to 6 years)", 3, 6, "F3C6D8", INK), ("Industry vs. inferiority\n(6 to 12 years)", 6, 12, "8E3B66", "FFFFFF")])]
    for name, bars in rows:
        label(s1, x, y, lw, 0.7, name, size=7, bold=True, color="FFFFFF", fillc=c)
        for t, a, b, fc, tc in bars:
            label(s1, ax(a) + 0.01, y, ax(b) - ax(a) - 0.02, 0.7, t, size=6.8, margin=0.015, bold=True, color=tc, fillc=fc, shape=MSO_SHAPE.RECTANGLE)
        y += 0.78
    # axis
    ln = s1.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ax0), Inches(y + 0.02), Inches(ax1), Inches(y + 0.02))
    ln.line.color.rgb = rgb("6B7280"); ln.line.width = Pt(0.75)
    for a in range(2, 13, 2):
        tb, tf = textbox(s1, ax(a) - 0.2, y + 0.03, 0.4, 0.16, margin=(0, 0))
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; add_runs(p, [(str(a), False, False, "4B5563")], 7)
    tb, tf = textbox(s1, x, y + 0.03, lw - 0.06, 0.16, margin=(0, 0))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT; add_runs(p, [("Age (years)", False, True, "4B5563")], 7)
    return y + 0.22 - h0

def person(x, y, s, color):
    box(s1, x, y, s * 0.55, s * 0.55, fillc=color, shape=MSO_SHAPE.OVAL)
    box(s1, x - s * 0.1, y + s * 0.6, s * 0.75, s * 0.6, fillc=color, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)

def v_play(x, y, w):
    """Figure 3: individual -> parallel -> cooperative play (Boyd and Bee, 2019)."""
    c = COL["socio"]; h0 = y
    n = 3; aw = 0.16; cw = (w - 2 * (aw + 0.06)) / n; ch = 1.25
    names = ["Individual play", "Parallel play", "Cooperative play\n(with common objectives)"]
    for k in range(n):
        cx0 = x + k * (cw + aw + 0.06)
        box(s1, cx0, y, cw, ch, fillc=TINT["socio"], line=c, lw=0.75, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        s = 0.26; py = y + 0.1
        if k == 0:
            person(cx0 + cw / 2 - 0.12, py, s, c)
            box(s1, cx0 + cw / 2 + 0.1, py + 0.22, 0.13, 0.13, fillc="F2A541", shape=MSO_SHAPE.RECTANGLE)
        elif k == 1:
            for dx in (-0.31, 0.1):
                person(cx0 + cw / 2 + dx, py, s, c)
                box(s1, cx0 + cw / 2 + dx + 0.19, py + 0.24, 0.08, 0.08, fillc="F2A541", shape=MSO_SHAPE.RECTANGLE)
        else:
            person(cx0 + cw / 2 - 0.3, py, s, c)
            person(cx0 + cw / 2 + 0.16, py, s, c)
            box(s1, cx0 + cw / 2 - 0.07, py + 0.2, 0.14, 0.14, fillc="F2A541", shape=MSO_SHAPE.OVAL)
        tb, tf = textbox(s1, cx0, y + 0.5, cw, 0.7, margin=(0.02, 0), anchor=MSO_ANCHOR.MIDDLE)
        for j, ln in enumerate(names[k].split("\n")):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
            add_runs(p, [(ln, j == 0, j == 1, c if j == 0 else INK)], 7.5)
        if k < n - 1:
            box(s1, cx0 + cw + 0.03, y + ch / 2 - 0.09, aw, 0.18, fillc=c, shape=MSO_SHAPE.RIGHT_ARROW)
    return ch

def v_bronf(x, y, w):
    """Figure 4: Bronfenbrenner's layers as concentric ellipses + chronosystem arrow."""
    c = COL["key"]; h = 2.2; h0 = y
    layers = [("macrosystem (culture, laws)", "C9DCEF"), ("exosystem (work of parents, local services)", "DCE8F5"),
              ("mesosystem (connections between them)", "EAF1F9"), ("microsystem (family, school)", "F7FAFD")]
    cxm, cym = x + w / 2, y + h / 2 - 0.12
    W0, H0 = w - 0.05, h - 0.3
    for k, (t, fc) in enumerate(layers):
        ww = W0 * (1 - 0.2 * k); hh = H0 * (1 - 0.2 * k)
        e = box(s1, cxm - ww / 2, cym - hh / 2, ww, hh, fillc=fc, line=c, lw=0.75, shape=MSO_SHAPE.OVAL)
        band = H0 * 0.1
        tb, tf = textbox(s1, cxm - 0.8, cym - hh / 2 + band / 2 - 0.08, 1.6, 0.16, margin=(0, 0), anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        nm = t.split(" (", 1)[0]
        add_runs(p, [(nm, True, False, c)], 7.5)
    ww = W0 * 0.2; hh = H0 * 0.2
    label(s1, cxm - 0.36, cym - 0.13, 0.72, 0.26, "The child", size=7.5, bold=True, color="FFFFFF", fillc=c, shape=MSO_SHAPE.OVAL)
    y2 = y + h - 0.3
    a = box(s1, x + 0.1, y2 + 0.04, w - 0.2, 0.24, fillc=c, shape=MSO_SHAPE.RIGHT_ARROW)
    tf = a.text_frame; tf.margin_top = tf.margin_bottom = Inches(0); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_runs(p, [("chronosystem", True, False, "FFFFFF"), (" (time)", False, False, "FFFFFF")], 7)
    return h

def v_chart(x, y, w):
    """Figure 5: Year 6 obesity by area deprivation, England 2024/25 (OHID, 2025)."""
    h = 1.45
    cd = CategoryChartData()
    cd.categories = ["Least deprived", "All Year 6", "Most deprived"]
    cd.add_series("Year 6 obesity (%)", (13.5, 22.2, 29.3))
    gf = s1.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = "Year 6 obesity chart"
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Year 6 obesity, England 2024/25 (%)"
    tp = ch.chart_title.text_frame.paragraphs[0]
    for r in tp.runs: r.font.size = Pt(8); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = rgb(INK)
    pl = ch.plots[0]; pl.gap_width = 45; pl.vary_by_categories = False
    pl.has_data_labels = True
    dl = pl.data_labels; dl.number_format = '0.0"%"'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(8); dl.font.bold = True; dl.font.name = FONT; dl.font.color.rgb = rgb(INK)
    ser = pl.series[0]
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb("7FA7CF")
    pt = ser.points[2]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(COL["key"])
    pt = ser.points[1]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb("4F7FB0")
    va = ch.value_axis; va.maximum_scale = 35; va.minimum_scale = 0; va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(7.5); ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = rgb(INK); ca.format.line.color.rgb = rgb("9CA3AF")
    return h

# ---------------- LAYOUT ----------------
GAP = 0.26  # vertical gap between panels
VG = 0.06  # gap between blocks in a panel
WTS = [1.17, 1.02, 0.86, 1.0, 0.95]
unit = (SW - 2 * M - (NC - 1) * G) / sum(WTS)
CWS = [w * unit for w in WTS]
cx = [M + sum(CWS[:k]) + k * G for k in range(NC)]
def span(c0, n): return cx[c0 + n - 1] + CWS[c0 + n - 1] - cx[c0]

def measure_panel(c0, ncols, key, blocks):
    col_h = []
    for ci, col in enumerate(blocks):
        sub_w = CWS[c0 + ci] - 2 * PAD
        hh = 0
        for b in col:
            if b[0] == "text": hh += measure_items(block_items(b[1], COL[key]), sub_w) + VG
            else: hh += b[4] + 0.22 + VG + 0.06
        col_h.append(hh)
    return HDR_H + 0.06 + max(col_h) + 0.04, col_h

def draw_panel(c0, y, ncols, key, title, blocks, pad_to=None):
    h_needed, col_h = measure_panel(c0, ncols, key, blocks)
    H = max(h_needed, pad_to or 0)
    panel(s1, cx[c0], y, span(c0, ncols), H, key, title)
    for ci, col in enumerate(blocks):
        bx = cx[c0 + ci] + PAD; sub_w = CWS[c0 + ci] - 2 * PAD
        by = y + HDR_H + 0.06
        for b in col:
            if b[0] == "text":
                by += place_text(b[1], key, bx, by, sub_w, name=f"{title} text") + VG
            else:
                _, fn, fig, cite, _hh = b
                by += 0.03
                used = fn(bx + 0.02, by, sub_w - 0.04)
                by += used + 0.03
                caption(s1, bx, by, sub_w, fig, cite, COL[key])
                by += 0.22 + VG
    report.append((title, round(y + H, 2), [round(h, 2) for h in col_h]))
    return H

VH = {"tl": 0.46 + 3 * 0.44, "st": 2 * 0.78 + 0.22, "pl": 1.25, "br": 2.2, "ch": 1.45}

H1 = draw_panel(0, TOP, 1, "intro", "1. Introduction", [[("text", intro)]])
y2 = TOP + H1 + GAP
H2 = draw_panel(0, y2, 1, "bio", "2. Biological Development",
                [[("text", bio), ("vis", v_timeline, 1, "(Boyd and Bee, 2019)", VH["tl"])]], pad_to=BOT - y2)
arrow(s1, cx[0] + CWS[0] / 2 - 0.13, TOP + H1 + 0.04, 0.26, GAP - 0.08, "down", COL["intro"])

H3 = draw_panel(1, TOP, 2, "psych", "3. Psychological Development",
                [[("text", psych_a)], [("vis", v_stages, 2, "(Boyd and Bee, 2019; Feldman, 2011)", VH["st"]), ("text", psych_b)]])
y4 = TOP + H3 + GAP
H4 = draw_panel(1, y4, 2, "socio", "4. Sociological Development",
                [[("text", socio[:6])], [("vis", v_play, 3, "(Boyd and Bee, 2019)", VH["pl"]), ("text", socio[6:])]], pad_to=BOT - y4)
mid = (cx[1] + span(1, 2) / 2)
arrow(s1, mid - 0.13, TOP + H3 + 0.04, 0.26, GAP - 0.08, "down", COL["psych"])
arrow(s1, cx[0] + CWS[0] + 0.03, TOP + HDR_H + 1.2, G - 0.06, 0.3, "right", COL["bio"])

key_blocks = [[("text", key_a), ("vis", v_bronf, 4, "(Boyd and Bee, 2019)", VH["br"]), ("text", key_b)],
              [("text", key_c), ("vis", v_chart, 5, "(OHID, 2025)", VH["ch"]), ("text", key_d)]]
H5 = draw_panel(3, TOP, 2, "key", "5. Key Influences", key_blocks)
y6 = TOP + H5 + GAP
H6 = draw_panel(3, y6, 2, "concl", "6. Conclusion", [[("text", concl[:3])], [("text", concl[3:])]], pad_to=BOT - y6)
arrow(s1, cx[2] + CWS[2] + 0.03, y4 + HDR_H + 0.5, G - 0.06, 0.3, "right", COL["socio"])
arrow(s1, cx[3] + span(3, 2) / 2 - 0.13, TOP + H5 + 0.04, 0.26, GAP - 0.08, "down", COL["key"])

print("BOT %.2f" % BOT)
for r in report: print(r)

# ======================= SLIDE 2 =======================
s2 = prs.slides.add_slide(BLANK)
bg = s2.background.fill; bg.solid(); bg.fore_color.rgb = rgb("FFFFFF")
tb, tf = textbox(s2, 0.5, 0.25, SW - 1, 0.6, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
add_runs(p, [("Annotated Bibliography", False, False, "000000")], 28)

REF_IDX = [61, 65, 69, 73, 76, 80, 84, 87, 91]
ANN_IDX = [63, 67, 71, 74, 78, 82, 85, 89, 93]
def ref_runs(i):
    out = []
    for r in DOC.paragraphs[i].runs:
        t = r.text
        for a, b in FIXES.get(i, []): t = t.replace(a, b)
        out.append((t, True, bool(r.italic), "000000", True))
    # strip leading/trailing whitespace on the paragraph
    if out:
        out[0] = (out[0][0].lstrip(),) + out[0][1:]
        out[-1] = (out[-1][0].rstrip(),) + out[-1][1:]
    return out

BIB = 10.5
LX, LW = 0.5, 7.6
RX, RW = 8.4, SW - 0.5 - 8.4
ty = 1.0
# column headings
for x_, w_, t_ in ((LX, LW, "Reference list"), (RX, RW, "Annotations")):
    tb, tf = textbox(s2, x_, ty, w_, 0.3, margin=(0.04, 0.02))
    add_runs(tf.paragraphs[0], [(t_, True, False, "000000")], 13)
ty += 0.36
ln = s2.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(LX), Inches(ty - 0.03), Inches(RX + RW), Inches(ty - 0.03))
ln.line.color.rgb = rgb("000000"); ln.line.width = Pt(0.75)

def url_safe_lines(runs, width_pt, size):
    # long URLs wrap at character level in PowerPoint; approximate by char-splitting long tokens
    rr = []
    for t, b, it, *_ in runs:
        parts = []
        for w in t.split(" "):
            while text_w(w, b, it, size) > width_pt:
                # cut
                k = len(w)
                while text_w(w[:k], b, it, size) > width_pt: k -= 1
                parts.append(w[:k]); w = w[k:]
            parts.append(w)
        rr.append((" ".join(parts), b, it))
    return count_lines(rr, width_pt, size)

rows_h = []
for ri, ai in zip(REF_IDX, ANN_IDX):
    rr = ref_runs(ri)
    hl = url_safe_lines(rr, (LW - 0.1) * 72, BIB) * BIB * LH
    ha = url_safe_lines([(T(ai), False, False)], (RW - 0.1) * 72, BIB) * BIB * LH
    h = max(hl, ha) / 72 + 0.07
    tb, tf = textbox(s2, LX, ty, LW, h, margin=(0.04, 0.02))
    p = tf.paragraphs[0]; p.line_spacing = 1.0
    add_runs(p, rr, BIB)
    tb, tf = textbox(s2, RX, ty, RW, h, margin=(0.04, 0.02))
    p = tf.paragraphs[0]; p.line_spacing = 1.0; p.alignment = PP_ALIGN.JUSTIFY
    add_runs(p, [(T(ai), False, False, "000000")], BIB)
    ty += h + 0.09
    rows_h.append(h)

ty += 0.02
tb, tf = textbox(s2, LX, ty, LW, 0.55, margin=(0.04, 0.02))
add_runs(tf.paragraphs[0], [("Image References", True, False, "000000")], 11)
p = tf.add_paragraph()
add_runs(p, [("[MISSING: image references – your content lists no photographs or illustrations. Add the full reference for any image you insert here (images are referenced only, not annotated).]", False, True, "B42318")], BIB)
print("slide 2 bottom: %.2f (limit %.2f)" % (ty + 0.55, SH - 0.3))

cp = prs.core_properties; cp.author = ""; cp.last_modified_by = ""; cp.title = "Human Development in Childhood"; cp.comments = ""
prs.save(OUT)
print("saved", OUT)
